from fastapi import FastAPI
from pydantic import BaseModel
from dotenv import load_dotenv
from fastapi.middleware.cors import CORSMiddleware
from groq import Groq

from docx import Document
from fastapi.responses import FileResponse
import os

load_dotenv()
app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

class ResumeData(BaseModel):
    experience: str
    jd: str

@app.post("/generate")
@app.post("/download-docx")
async def generate_resume(data: ResumeData):
    try:
        client = Groq(api_key=os.getenv("GROQ_API_KEY"))
        prompt = f"Act as a career expert. Optimize these resume points for this job description: {data.jd}\n\nExperience: {data.experience}"
        
        chat_completion = client.chat.completions.create(
            messages=[{"role": "user", "content": prompt}],
            model="llama-3.3-70b-versatile", 
        )
        return {"improved_resume": chat_completion.choices[0].message.content}
    except Exception as e:
        return {"error": str(e)}
    
async def generate_docx(data: ResumeData):
    doc = Document()
    doc.add_heading('Resume', 0)
    doc.add_paragraph(data.improved_resume) # ဒီနေရာမှာ AI ရလဒ်ထည့်
    doc.save("resume.docx")
    return FileResponse("resume.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")